from __future__ import annotations

import re
import zipfile
from datetime import time
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from pydantic import ValidationError

from personal_assistant.schemas import CourseInput

ALIASES: dict[str, tuple[str, ...]] = {
    "course_name": ("课程名称", "课程名", "课程", "科目", "course_name", "course"),
    "weekday": ("星期", "周几", "上课星期", "星期几", "weekday", "day"),
    "start_period": ("开始节次", "起始节次", "节次", "上课节次", "start_period"),
    "end_period": ("结束节次", "end_period"),
    "start_time": ("开始时间", "上课时间", "start_time"),
    "end_time": ("结束时间", "end_time"),
    "weeks": ("周次", "上课周次", "教学周", "weeks"),
    "week_parity": ("单双周", "周类型", "week_parity"),
    "location": ("地点", "教室", "上课地点", "location", "room"),
    "teacher": ("教师", "老师", "任课教师", "teacher"),
    "notes": ("备注", "notes"),
}

FIELD_LABELS = {
    "course_name": "课程名称（必填）",
    "weekday": "星期（必填，周一为1）",
    "start_period": "开始节次",
    "end_period": "结束节次",
    "start_time": "开始时间",
    "end_time": "结束时间",
    "weeks": "周次（如1-16周）",
    "week_parity": "单双周（单周/双周）",
    "location": "地点",
    "teacher": "教师",
    "notes": "备注",
}

WEEKDAY_HEADERS = {
    "星期一": 1,
    "周一": 1,
    "星期二": 2,
    "周二": 2,
    "星期三": 3,
    "周三": 3,
    "星期四": 4,
    "周四": 4,
    "星期五": 5,
    "周五": 5,
    "星期六": 6,
    "周六": 6,
    "星期日": 7,
    "星期天": 7,
    "周日": 7,
    "周天": 7,
}

MAX_XLSX_ENTRIES = 2_048
MAX_XLSX_EXPANDED_BYTES = 64 * 1024 * 1024
MAX_XLSX_ENTRY_BYTES = 32 * 1024 * 1024
MAX_XLSX_COMPRESSION_RATIO = 200


def validate_xlsx_archive(data: bytes) -> None:
    """Reject malformed or unusually large OOXML archives before openpyxl parses them."""
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_XLSX_ENTRIES:
                raise ValueError("工作簿包含过多文件条目")
            expanded_bytes = 0
            for entry in entries:
                if entry.flag_bits & 0x1:
                    raise ValueError("不支持加密的 Excel 工作簿")
                if entry.file_size > MAX_XLSX_ENTRY_BYTES:
                    raise ValueError("工作簿中的单个文件解压后过大")
                expanded_bytes += entry.file_size
                if expanded_bytes > MAX_XLSX_EXPANDED_BYTES:
                    raise ValueError("工作簿解压后的总大小超过限制")
                if entry.file_size and not entry.compress_size:
                    raise ValueError("工作簿压缩条目格式异常")
                if (
                    entry.compress_size
                    and entry.file_size / entry.compress_size > MAX_XLSX_COMPRESSION_RATIO
                ):
                    raise ValueError("工作簿压缩率异常，拒绝解析")
    except zipfile.BadZipFile as exc:
        raise ValueError("文件不是有效的 .xlsx 工作簿") from exc


def period_time(period: int | None) -> str | None:
    times = {
        1: ("08:00", "08:45"),
        2: ("08:50", "09:35"),
        3: ("09:50", "10:35"),
        4: ("10:40", "11:25"),
        5: ("11:30", "12:15"),
        6: ("14:00", "14:45"),
        7: ("14:50", "15:35"),
        8: ("15:50", "16:35"),
        9: ("16:40", "17:25"),
        10: ("17:30", "18:15"),
        11: ("19:00", "19:45"),
        12: ("19:50", "20:35"),
        13: ("20:40", "21:25"),
        14: ("21:30", "22:15"),
    }
    return times.get(period, (None, None))[0] if period else None


def period_range(start_period: int | None, end_period: int | None) -> tuple[str | None, str | None]:
    starts = {
        1: "08:00",
        2: "08:50",
        3: "09:50",
        4: "10:40",
        5: "11:30",
        6: "14:00",
        7: "14:50",
        8: "15:50",
        9: "16:40",
        10: "17:30",
        11: "19:00",
        12: "19:50",
        13: "20:40",
        14: "21:30",
    }
    ends = {
        1: "08:45",
        2: "09:35",
        3: "10:35",
        4: "11:25",
        5: "12:15",
        6: "14:45",
        7: "15:35",
        8: "16:35",
        9: "17:25",
        10: "18:15",
        11: "19:45",
        12: "20:35",
        13: "21:25",
        14: "22:15",
    }
    return starts.get(start_period), ends.get(end_period)


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, time):
        return value.strftime("%H:%M")
    return str(value).strip()


def _normal(value: Any) -> str:
    return re.sub(r"[\s_（）()\-]+", "", _text(value)).lower()


def suggest_mapping(headers: list[Any]) -> dict[str, str]:
    normalized = {_normal(header): _text(header) for header in headers if header is not None}
    result: dict[str, str] = {}
    for field, aliases in ALIASES.items():
        candidates = {_normal(alias) for alias in aliases}
        for header, original in normalized.items():
            if header in candidates:
                result[field] = original
                break
    return result


def workbook_info(data: bytes, sheet_name: str | None, header_row: int) -> dict[str, Any]:
    validate_xlsx_archive(data)
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        if sheet_name and sheet_name not in workbook.sheetnames:
            raise ValueError(f"找不到工作表：{sheet_name}")
        selected = sheet_name or workbook.sheetnames[0]
        sheet = workbook[selected]
        grid_header = _find_weekday_header(sheet)
        if grid_header is not None:
            return {
                "sheets": workbook.sheetnames,
                "sheet_name": selected,
                "header_row": grid_header[0],
                "headers": ["星期", "课程安排（合并单元格）"],
                "suggested_mapping": {},
                "sample_rows": [],
                "format": "weekday_grid",
            }
        headers = [_text(cell.value) for cell in sheet[header_row]]
        while headers and not headers[-1]:
            headers.pop()
        sample = []
        for row in sheet.iter_rows(
            min_row=header_row + 1, max_row=header_row + 8, values_only=True
        ):
            if any(value is not None for value in row):
                sample.append([_text(value) for value in row[: len(headers)]])
        return {
            "sheets": workbook.sheetnames,
            "sheet_name": selected,
            "header_row": header_row,
            "headers": headers,
            "suggested_mapping": suggest_mapping(headers),
            "sample_rows": sample,
            "format": "row_table",
        }
    finally:
        workbook.close()


def _find_weekday_header(sheet) -> tuple[int, dict[int, int]] | None:
    max_rows = min(sheet.max_row or 0, 20)
    for row_index, row in enumerate(
        sheet.iter_rows(min_row=1, max_row=max_rows, values_only=True), start=1
    ):
        found: dict[int, int] = {}
        for column, value in enumerate(row, start=1):
            normalized = _text(value).replace(" ", "")
            if normalized in WEEKDAY_HEADERS and column >= 3:
                found[column] = WEEKDAY_HEADERS[normalized]
        if len(set(found.values())) >= 5:
            return row_index, found
    return None


def _parse_grid_cell(value: Any, weekday: int) -> tuple[list[dict[str, Any]], list[str]]:
    text = _text(value)
    starts = list(re.finditer(r"[（(]\s*本\s*[）)]", text))
    if not starts:
        return [], []
    courses: list[dict[str, Any]] = []
    errors: list[str] = []
    for index, marker in enumerate(starts):
        block = text[
            marker.end() : starts[index + 1].start() if index + 1 < len(starts) else len(text)
        ]
        lines = [line.strip() for line in block.splitlines() if line.strip()]
        if not lines:
            continue
        course_name = re.split(r"\s{2,}", lines[0], maxsplit=1)[0].strip()
        period_match = re.search(r"第\s*(\d+)\s*[-~—至到,，、]\s*(\d+)\s*节", block)
        if not period_match:
            period_match = re.search(r"第\s*(\d+)\s*节", block)
        if not course_name or not period_match:
            errors.append(f"课程块“{course_name or lines[0]}”缺少课程名或节次")
            continue
        start_period = int(period_match.group(1))
        end_period = (
            int(period_match.group(2))
            if period_match.lastindex and period_match.lastindex >= 2
            else start_period
        )
        teacher_entries = re.findall(r"([^,，\n\[\]]+?)\s*\[([^\]]+)\]", block)
        teachers: list[str] = []
        weeks: set[int] = set()
        for teacher_name, weeks_text in teacher_entries:
            teacher_name = teacher_name.strip()
            if teacher_name and teacher_name not in teachers:
                teachers.append(teacher_name)
            weeks.update(_weeks(weeks_text))
        location = ""
        for line in lines[1:]:
            if "[" in line or re.search(r"第\s*\d+.*节", line) or "周" in line:
                continue
            if re.search(r"（本）|\(本\)", line):
                continue
            location = line
            break
        if start_period < 1 or end_period > 14 or end_period < start_period:
            errors.append(f"课程块“{course_name}”节次范围无效：{start_period}-{end_period}")
            continue
        start_time, end_time = period_range(start_period, end_period)
        courses.append(
            {
                "course_name": course_name,
                "weekday": weekday,
                "start_period": start_period,
                "end_period": end_period,
                "start_time": start_time,
                "end_time": end_time,
                "weeks": sorted(weeks),
                "week_parity": "all",
                "location": location,
                "teacher": ", ".join(teachers),
                "notes": "",
            }
        )
    return courses, errors


def parse_weekday_grid_courses(data: bytes, sheet_name: str | None) -> dict[str, Any]:
    validate_xlsx_archive(data)
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        if sheet_name and sheet_name not in workbook.sheetnames:
            raise ValueError(f"找不到工作表：{sheet_name}")
        selected = sheet_name or workbook.sheetnames[0]
        sheet = workbook[selected]
        header = _find_weekday_header(sheet)
        if header is None:
            raise ValueError("没有识别到星期一至星期日的表头")
        header_row, day_columns = header
        parsed: list[dict[str, Any]] = []
        errors: list[dict[str, Any]] = []
        seen: set[tuple[Any, ...]] = set()
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1
        ):
            for column, weekday in day_columns.items():
                if column > len(row) or row[column - 1] is None:
                    continue
                items, cell_errors = _parse_grid_cell(row[column - 1], weekday)
                for course in items:
                    key = tuple(
                        (field, tuple(value) if isinstance(value, list) else value)
                        for field, value in course.items()
                    )
                    if key not in seen:
                        parsed.append(course)
                        seen.add(key)
                for message in cell_errors:
                    errors.append({"row": row_number, "column": column, "error": message})
        if not parsed and not errors:
            errors.append({"row": header_row, "column": 0, "error": "课表中没有识别到课程"})
        return {
            "sheet_name": selected,
            "headers": ["星期", "课程安排（合并单元格）"],
            "mapping": {},
            "courses": parsed,
            "errors": errors,
            "format": "weekday_grid",
        }
    finally:
        workbook.close()


def _weekday(value: Any) -> int:
    text = _text(value).lower()
    names = {
        "一": 1,
        "二": 2,
        "三": 3,
        "四": 4,
        "五": 5,
        "六": 6,
        "日": 7,
        "天": 7,
        "mon": 1,
        "tue": 2,
        "wed": 3,
        "thu": 4,
        "fri": 5,
        "sat": 6,
        "sun": 7,
    }
    for name, number in names.items():
        if name in text:
            return number
    match = re.search(r"\d+", text)
    if match:
        number = int(match.group())
        if 1 <= number <= 7:
            return number
    raise ValueError(f"无法识别星期：{text or '空值'}")


def _period(value: Any) -> int | None:
    text = _text(value)
    match = re.search(r"\d+", text)
    return int(match.group()) if match else None


def _period_end(value: Any) -> int | None:
    matches = re.findall(r"\d+", _text(value))
    return int(matches[1]) if len(matches) >= 2 else None


def _time_value(value: Any) -> str | None:
    if value is None or _text(value) == "":
        return None
    if isinstance(value, time):
        return value.strftime("%H:%M")
    text = _text(value)
    match = re.search(r"(\d{1,2}):(\d{2})", text)
    if not match:
        return None
    hour, minute = int(match.group(1)), int(match.group(2))
    if hour > 23 or minute > 59:
        return None
    return f"{hour:02d}:{minute:02d}"


def _weeks(value: Any) -> list[int]:
    text = _text(value)
    if not text:
        return []
    result: set[int] = set()
    for start, end in re.findall(r"(\d+)\s*[-~—至到]\s*(\d+)", text):
        low, high = int(start), int(end)
        if low <= high <= 60:
            result.update(range(low, high + 1))
    covered = re.sub(r"\d+\s*[-~—至到]\s*\d+", "", text)
    for value_text in re.findall(r"\d+", covered):
        number = int(value_text)
        if 1 <= number <= 60:
            result.add(number)
    return sorted(result)


def _parity(value: Any, weeks_text: Any) -> str:
    text = f"{_text(value)} {_text(weeks_text)}".lower()
    if "单" in text or "odd" in text:
        return "odd"
    if "双" in text or "even" in text:
        return "even"
    return "all"


def parse_workbook_courses(
    data: bytes,
    sheet_name: str | None,
    header_row: int,
    mapping: dict[str, str],
) -> dict[str, Any]:
    validate_xlsx_archive(data)
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        if sheet_name and sheet_name not in workbook.sheetnames:
            raise ValueError(f"找不到工作表：{sheet_name}")
        selected = sheet_name or workbook.sheetnames[0]
        sheet = workbook[selected]
        headers = [_text(cell.value) for cell in sheet[header_row]]
        while headers and not headers[-1]:
            headers.pop()
        positions = {header: index for index, header in enumerate(headers) if header}
        required = [field for field in ("course_name", "weekday") if not mapping.get(field)]
        if required:
            raise ValueError("请先指定列映射：" + ", ".join(required))
        missing_headers = [value for value in mapping.values() if value and value not in positions]
        if missing_headers:
            raise ValueError("列映射中的表头不存在：" + ", ".join(missing_headers))

        def get_value(row: tuple[Any, ...], field: str) -> Any:
            header = mapping.get(field)
            return (
                row[positions[header]]
                if header in positions and positions[header] < len(row)
                else None
            )

        parsed: list[dict[str, Any]] = []
        errors: list[dict[str, Any]] = []
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1
        ):
            if not any(value is not None for value in row):
                continue
            try:
                week_text = get_value(row, "weeks")
                start_period_raw = get_value(row, "start_period")
                end_period_raw = get_value(row, "end_period")
                end_period = _period(end_period_raw)
                if end_period is None:
                    end_period = _period_end(start_period_raw)
                course = CourseInput(
                    course_name=_text(get_value(row, "course_name")),
                    weekday=_weekday(get_value(row, "weekday")),
                    start_period=_period(start_period_raw),
                    end_period=end_period,
                    start_time=_time_value(get_value(row, "start_time")),
                    end_time=_time_value(get_value(row, "end_time")),
                    weeks=_weeks(week_text),
                    week_parity=_parity(get_value(row, "week_parity"), week_text),
                    location=_text(get_value(row, "location")),
                    teacher=_text(get_value(row, "teacher")),
                    notes=_text(get_value(row, "notes")),
                )
                if not course.start_period and not course.start_time:
                    raise ValueError("需要开始节次或开始时间")
                parsed.append(course.model_dump(mode="json"))
            except (ValueError, ValidationError) as exc:
                errors.append({"row": row_number, "error": str(exc)})
        return {
            "sheet_name": selected,
            "headers": headers,
            "mapping": mapping,
            "courses": parsed,
            "errors": errors,
        }
    finally:
        workbook.close()
